"""Partie 3.1.9 -- tests for api/services/headings_extraction.py."""

import tempfile
from pathlib import Path

from api.services.headings_extraction import (
    build_section_hierarchy,
    extract_headings_docx,
    extract_headings_generic,
    extract_headings_html,
    extract_headings_markdown,
    extract_headings_pdf,
    get_section_context,
    get_section_path,
    split_by_headings,
)

_MARKDOWN = "# Title\n\nIntro text.\n\n## Section A\n\nContent A.\n\n### Sub A.1\n\nDeep content.\n\n## Section B\n\nContent B.\n"


def test_extract_headings_markdown_finds_real_headings_in_order():
    """Validation criterion: l'extraction des titres Markdown fonctionne."""
    headings = extract_headings_markdown(_MARKDOWN)
    assert [(h["level"], h["text"]) for h in headings] == [
        (1, "Title"), (2, "Section A"), (3, "Sub A.1"), (2, "Section B"),
    ]


def test_extract_headings_markdown_ignores_hash_inside_a_code_fence():
    md = "# Real Heading\n\n```\n# not a heading, inside a fence\n```\n\nMore text.\n"
    headings = extract_headings_markdown(md)
    assert [h["text"] for h in headings] == ["Real Heading"]


def test_extract_headings_markdown_position_points_at_the_real_title_text():
    """A regression test for a real bug found while building this
    étape: position must land on the title itself, not the whole line
    (including the real `#` marker)."""
    headings = extract_headings_markdown(_MARKDOWN)
    for heading in headings:
        assert _MARKDOWN[heading["position"]:heading["position"] + len(heading["text"])] == heading["text"]


def test_extract_headings_markdown_is_empty_for_text_with_no_headings():
    assert extract_headings_markdown("Just a plain paragraph, no headings at all.") == []


def test_extract_headings_html_finds_real_h1_to_h6_tags():
    """Validation criterion: l'extraction des titres HTML fonctionne."""
    html = "<h1>Main</h1><p>intro</p><h2>Sub</h2><p>more</p><h3>Deeper</h3>"
    headings = extract_headings_html(html)
    assert [(h["level"], h["text"]) for h in headings] == [(1, "Main"), (2, "Sub"), (3, "Deeper")]


def test_extract_headings_generic_detects_numbered_outline():
    text = "1. Introduction\nSome text.\n1.1 Background\nMore text.\n1.1.1 Detail\nEven more."
    headings = extract_headings_generic(text)
    assert [(h["level"], h["text"]) for h in headings] == [(1, "Introduction"), (2, "Background"), (3, "Detail")]


def test_extract_headings_generic_detects_all_caps_lines():
    text = "INTRODUCTION\nSome real body text here."
    headings = extract_headings_generic(text)
    assert headings == [{"level": 1, "text": "INTRODUCTION", "position": 0}]


def test_extract_headings_generic_position_points_at_the_real_title_text():
    text = "1.1 Background\nSome text."
    headings = extract_headings_generic(text)
    heading = headings[0]
    assert text[heading["position"]:heading["position"] + len(heading["text"])] == heading["text"]


def _real_pdf_with_heading() -> str:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Chapter One", fontsize=22)
    page.insert_text((72, 110), "A real, ordinary body paragraph with plenty of real words in it.", fontsize=11)
    path = tempfile.mktemp(suffix=".pdf")
    doc.save(path)
    doc.close()
    return path


def test_extract_headings_pdf_finds_a_real_larger_span():
    """Validation criterion: l'extraction des titres PDF fonctionne."""
    path = _real_pdf_with_heading()
    try:
        headings = extract_headings_pdf(path)
        assert [h["text"] for h in headings] == ["Chapter One"]
        assert headings[0]["level"] == 1
    finally:
        Path(path).unlink(missing_ok=True)


def test_extract_headings_pdf_is_empty_for_uniform_font_size():
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "First line of uniform text.", fontsize=11)
    page.insert_text((72, 100), "Second line of uniform text.", fontsize=11)
    path = tempfile.mktemp(suffix=".pdf")
    doc.save(path)
    doc.close()
    try:
        assert extract_headings_pdf(path) == []
    finally:
        Path(path).unlink(missing_ok=True)


def _real_docx_with_heading() -> str:
    import docx

    document = docx.Document()
    document.add_heading("Chapter One", level=1)
    document.add_paragraph("A real, ordinary body paragraph.")
    document.add_heading("Section 1.1", level=2)
    path = tempfile.mktemp(suffix=".docx")
    document.save(path)
    return path


def test_extract_headings_docx_finds_real_heading_styles():
    """Validation criterion: l'extraction des titres DOCX fonctionne."""
    path = _real_docx_with_heading()
    try:
        headings = extract_headings_docx(path)
        assert [(h["level"], h["text"]) for h in headings] == [(1, "Chapter One"), (2, "Section 1.1")]
    finally:
        Path(path).unlink(missing_ok=True)


def test_split_by_headings_returns_real_clean_section_content():
    headings = extract_headings_markdown(_MARKDOWN)
    sections = split_by_headings(_MARKDOWN, headings)
    assert sections[0] == "Intro text."
    assert sections[1] == "Content A."
    assert sections[-1] == "Content B."
    # A real regression check for a bug found while building this étape:
    # a section's own trailing content must never leak the NEXT
    # heading's own real markup prefix (e.g. "##").
    for section in sections:
        assert "#" not in section


def test_build_section_hierarchy_nests_real_headings_correctly():
    """Validation criterion: la hiérarchie des sections est correcte."""
    headings = extract_headings_markdown(_MARKDOWN)
    tree = build_section_hierarchy(headings, _MARKDOWN)
    assert len(tree) == 1
    assert tree[0]["text"] == "Title"
    assert [child["text"] for child in tree[0]["children"]] == ["Section A", "Section B"]
    assert tree[0]["children"][0]["children"][0]["text"] == "Sub A.1"


def test_get_section_context_finds_the_real_enclosing_heading():
    headings = extract_headings_markdown(_MARKDOWN)
    deep_position = _MARKDOWN.index("Deep content")
    context = get_section_context(_MARKDOWN, deep_position, headings)
    assert context["text"] == "Sub A.1"


def test_get_section_context_is_none_before_the_first_heading():
    headings = extract_headings_markdown(_MARKDOWN)
    assert get_section_context(_MARKDOWN, 0, headings) is None


def test_get_section_path_returns_the_real_full_breadcrumb():
    headings = extract_headings_markdown(_MARKDOWN)
    deep_position = _MARKDOWN.index("Deep content")
    assert get_section_path(headings, deep_position) == ["Title", "Section A", "Sub A.1"]


def test_get_section_path_drops_stale_deeper_levels_after_a_sibling():
    headings = extract_headings_markdown(_MARKDOWN)
    end_position = len(_MARKDOWN)
    # "Section B" (level 2) comes after "Sub A.1" (level 3) -- the
    # stale level-3 entry must be dropped, not kept alongside it.
    assert get_section_path(headings, end_position) == ["Title", "Section B"]
