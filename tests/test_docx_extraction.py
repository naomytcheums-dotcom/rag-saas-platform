"""
Partie 2.1.2 -- real DOCX extraction tests (api/services/docx_extraction.py).
Same "no mocking, generate and extract for real" discipline as
tests/test_pdf_extraction.py: python-docx both BUILDS the test DOCX
files (real headings/tables/properties) and EXTRACTS from them.
"""

import io
import zipfile

import docx
import pytest

from api.services.docx_extraction import extract_docx_footnotes, extract_docx_metadata, extract_docx_styles, extract_docx_tables, extract_docx_text


@pytest.fixture
def real_docx_path(tmp_path):
    """A real DOCX: a heading, two body paragraphs, a subheading, and a
    real 3x2 table -- built with python-docx itself, saved to a real
    temp file (these functions take a file PATH, this step's own
    literal signature)."""
    document = docx.Document()
    document.core_properties.author = "pytest"
    document.core_properties.title = "Real Test DOCX"
    document.core_properties.keywords = "alpha; beta; gamma"

    document.add_heading("Main Title", level=1)
    document.add_paragraph("This is real body text under the main title.")
    document.add_heading("Subsection", level=2)
    document.add_paragraph("Real body text under the subsection.")

    table = document.add_table(rows=3, cols=2)
    data = [["Name", "Age"], ["Alice", "30"], ["Bob", "25"]]
    for r in range(3):
        for c in range(2):
            table.cell(r, c).text = data[r][c]

    path = tmp_path / "real_test.docx"
    document.save(str(path))
    return str(path)


@pytest.fixture
def corrupt_docx_path(tmp_path):
    path = tmp_path / "corrupt.docx"
    path.write_bytes(b"this is not a real docx, just garbage bytes")
    return str(path)


@pytest.fixture
def malformed_xml_docx_path(tmp_path):
    """A real, structurally valid ZIP with the right `word/document.xml`
    part present, but genuinely malformed XML inside it -- confirmed
    for real (see api/services/docx_extraction.py's own module
    docstring) to be a DIFFERENT, equally real corruption case from
    corrupt_docx_path above: this one raises a bare AttributeError from
    deep inside python-docx's own object model, not any DOCX-specific
    exception, which is exactly why _open catches broadly."""
    path = tmp_path / "malformed.docx"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", b"not valid xml at all <<<")
        archive.writestr("[Content_Types].xml", b"<Types/>")
    return str(path)


# ------------------------------------------------------------------- text --

def test_extract_docx_text_returns_every_real_paragraph(real_docx_path):
    """Validation criterion: text extraction works."""
    text = extract_docx_text(real_docx_path)
    assert "Main Title" in text
    assert "This is real body text under the main title." in text
    assert "Subsection" in text
    assert "Real body text under the subsection." in text


def test_extract_docx_text_raises_a_clear_error_for_a_corrupt_docx(corrupt_docx_path):
    """Vision critique Q2 -- what happens with a corrupted DOCX: a
    clear ValueError, not an unhandled python-docx-specific exception."""
    with pytest.raises(ValueError):
        extract_docx_text(corrupt_docx_path)


def test_extract_docx_text_raises_a_clear_error_for_malformed_internal_xml(malformed_xml_docx_path):
    """The SECOND, real corruption case this module's own docstring
    documents -- a valid ZIP/OPC package but broken XML inside it,
    confirmed to raise a bare AttributeError from python-docx itself if
    not caught broadly."""
    with pytest.raises(ValueError):
        extract_docx_text(malformed_xml_docx_path)


# ----------------------------------------------------------------- tables --

def test_extract_docx_tables_finds_the_real_table_with_header_row(real_docx_path):
    """Validation criterion: table extraction works. The DataFrame's
    own column headers come from the table's real first row."""
    tables = extract_docx_tables(real_docx_path)
    assert len(tables) == 1
    df = tables[0]
    assert list(df.columns) == ["Name", "Age"]
    assert df.iloc[0].tolist() == ["Alice", "30"]
    assert df.iloc[1].tolist() == ["Bob", "25"]


def test_extract_docx_tables_returns_empty_list_when_there_are_none(tmp_path):
    document = docx.Document()
    document.add_paragraph("Just plain text, no tables here.")
    path = tmp_path / "no_tables.docx"
    document.save(str(path))

    assert extract_docx_tables(str(path)) == []


def test_extract_docx_tables_raises_for_a_corrupt_docx(corrupt_docx_path):
    with pytest.raises(ValueError):
        extract_docx_tables(corrupt_docx_path)


# --------------------------------------------------------------- metadata --

def test_extract_docx_metadata_returns_the_real_title_author_and_paragraph_count(real_docx_path):
    """Validation criterion: metadata is extracted (author, title, etc.)."""
    metadata = extract_docx_metadata(real_docx_path)
    assert metadata["title"] == "Real Test DOCX"
    assert metadata["author"] == "pytest"
    assert metadata["paragraph_count"] == 4  # 2 headings + 2 body paragraphs


def test_extract_docx_metadata_returns_the_real_keywords_tag(real_docx_path):
    """Partie 2.2.5 -- the real OOXML "Tags" core property."""
    metadata = extract_docx_metadata(real_docx_path)
    assert metadata["keywords"] == "alpha; beta; gamma"


def test_extract_docx_metadata_raises_for_a_corrupt_docx(corrupt_docx_path):
    with pytest.raises(ValueError):
        extract_docx_metadata(corrupt_docx_path)


# ----------------------------------------------------------------- styles --

def test_extract_docx_styles_identifies_real_headings_and_body_text():
    """Validation criterion: styles are extracted (headings, bold,
    etc.) -- confirmed against python-docx's own real style names, not
    a guess at what they'd be called."""
    document = docx.Document()
    document.add_heading("A Heading", level=1)
    document.add_paragraph("Regular body text.")
    path_holder = io.BytesIO()
    document.save(path_holder)

    import tempfile
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp.write(path_holder.getvalue())
        tmp_path = tmp.name

    styles = extract_docx_styles(tmp_path)
    assert styles == [
        {"text": "A Heading", "style": "Heading 1"},
        {"text": "Regular body text.", "style": "Normal"},
    ]


def test_extract_docx_styles_skips_empty_paragraphs(tmp_path):
    document = docx.Document()
    document.add_paragraph("Real text.")
    document.add_paragraph("")  # an empty paragraph Word itself would still store
    path = tmp_path / "with_empty_paragraph.docx"
    document.save(str(path))

    styles = extract_docx_styles(str(path))
    assert len(styles) == 1
    assert styles[0]["text"] == "Real text."


def test_extract_docx_styles_raises_for_a_corrupt_docx(corrupt_docx_path):
    with pytest.raises(ValueError):
        extract_docx_styles(corrupt_docx_path)


# --------------------------------------------------- footnotes (Partie 3.1.3) --

def _docx_bytes_with_real_footnote() -> bytes:
    """python-docx has no public API to ADD a footnote either (the
    same real gap this étape's own `extract_docx_footnotes` was built
    to read around) -- a real, minimal, hand-crafted OOXML
    `word/footnotes.xml` part is injected directly into a real,
    otherwise normal python-docx-generated package, the same real
    "manipulate the real ZIP directly" technique
    tests/test_documents_integration.py's own corrupt-DOCX test already
    uses. Includes the two real, structural separator footnotes Word
    itself always emits (ids -1/0, never real content) alongside one
    real footnote with real text, so the test also proves those two
    are correctly excluded."""
    import io
    import zipfile

    import docx

    document = docx.Document()
    document.add_paragraph("Body text with a real footnote reference.")
    base_buffer = io.BytesIO()
    document.save(base_buffer)
    base_buffer.seek(0)

    footnotes_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
        '<w:footnote w:id="1"><w:p><w:r><w:t>This is a real footnote with real content.</w:t></w:r></w:p></w:footnote>'
        "</w:footnotes>"
    )

    output_buffer = io.BytesIO()
    with zipfile.ZipFile(base_buffer, "r") as zip_in, zipfile.ZipFile(output_buffer, "w") as zip_out:
        for item in zip_in.infolist():
            data = zip_in.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(
                    b"</Types>",
                    b'<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>',
                )
            elif item.filename == "word/_rels/document.xml.rels":
                data = data.replace(
                    b"</Relationships>",
                    b'<Relationship Id="rIdFootnotes" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>',
                )
            zip_out.writestr(item, data)
        zip_out.writestr("word/footnotes.xml", footnotes_xml)
    return output_buffer.getvalue()


def test_extract_docx_footnotes_finds_real_footnote_content(tmp_path):
    """Validation criterion (Partie 3.1.3) -- un vrai gap trouvé en
    révisant les extracteurs existants : python-docx n'a aucune API
    publique pour les notes de bas de page."""
    path = tmp_path / "with_footnote.docx"
    path.write_bytes(_docx_bytes_with_real_footnote())

    footnotes = extract_docx_footnotes(str(path))
    assert footnotes == ["This is a real footnote with real content."]


def test_extract_docx_footnotes_excludes_real_structural_separators(tmp_path):
    """The two real, structural separator footnotes (ids -1/0) Word
    itself always emits are never real content."""
    path = tmp_path / "with_footnote.docx"
    path.write_bytes(_docx_bytes_with_real_footnote())

    footnotes = extract_docx_footnotes(str(path))
    assert len(footnotes) == 1  # not 3 -- the two structural ones excluded


def test_extract_docx_footnotes_returns_empty_list_for_a_real_docx_with_no_footnotes_part(real_docx_path):
    assert extract_docx_footnotes(real_docx_path) == []


def test_extract_docx_footnotes_raises_for_a_corrupt_docx(corrupt_docx_path):
    with pytest.raises(ValueError):
        extract_docx_footnotes(corrupt_docx_path)
