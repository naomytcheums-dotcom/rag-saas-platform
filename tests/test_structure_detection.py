"""Partie 3.1.8 -- tests for api/services/structure_detection.py."""

import tempfile
from pathlib import Path

from api.services.structure_detection import (
    StructureElement,
    detect_structure_docx,
    detect_structure_html,
    detect_structure_markdown,
    detect_structure_pdf,
    detect_structure_text,
    structure_to_json,
    structure_to_markdown,
)

_MARKDOWN = (
    "# Title\n\nIntro paragraph.\n\n- item one\n- item two\n\n"
    "## Section A\n\n```python\ndef f():\n    pass\n```\n\nSome text.\n"
)


def test_detect_structure_markdown_finds_every_real_element_type():
    """Validation criterion: la détection de structure Markdown fonctionne."""
    structure = detect_structure_markdown(_MARKDOWN)
    types = [e.type for e in structure]
    assert types == ["heading", "paragraph", "list_item", "list_item", "heading", "code_block", "paragraph"]
    assert structure[0].level == 1
    assert structure[4].level == 2
    assert "def f()" in structure[5].content


def test_detect_structure_markdown_is_empty_for_empty_text():
    assert detect_structure_markdown("") == []


def test_detect_structure_html_finds_headings_paragraphs_and_lists():
    """Validation criterion: la détection de structure HTML fonctionne."""
    html = "<h1>Main</h1><p>intro</p><ul><li>a</li><li>b</li></ul><table><tr><td>x</td></tr></table>"
    structure = detect_structure_html(html)
    types = [e.type for e in structure]
    assert types == ["heading", "paragraph", "list_item", "list_item", "table"]
    assert structure[0].level == 1


def test_detect_structure_text_finds_headings_and_groups_paragraphs():
    text = "INTRODUCTION\nSome text here.\nMore text.\n\n1. Section One\nContent one."
    structure = detect_structure_text(text)
    assert [(e.type, e.level, e.content) for e in structure] == [
        ("heading", 1, "INTRODUCTION"),
        ("paragraph", None, "Some text here. More text."),
        ("heading", 1, "Section One"),
        ("paragraph", None, "Content one."),
    ]


def _real_pdf_with_heading_and_body() -> str:
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Chapter One", fontsize=22)
    page.insert_text((72, 110), "A real, ordinary body paragraph with plenty of real words in it.", fontsize=11)
    path = tempfile.mktemp(suffix=".pdf")
    doc.save(path)
    doc.close()
    return path


def test_detect_structure_pdf_distinguishes_heading_from_body():
    """Validation criterion: la détection de structure PDF fonctionne."""
    path = _real_pdf_with_heading_and_body()
    try:
        structure = detect_structure_pdf(path)
        assert structure[0].type == "heading"
        assert structure[0].content == "Chapter One"
        assert structure[1].type == "paragraph"
        assert "body paragraph" in structure[1].content
    finally:
        Path(path).unlink(missing_ok=True)


def _real_docx_with_heading_and_table() -> str:
    import docx

    document = docx.Document()
    document.add_heading("Chapter One", level=1)
    document.add_paragraph("A real body paragraph.")
    table = document.add_table(rows=2, cols=2)
    data = [["Name", "Value"], ["real", "table"]]
    for r in range(2):
        for c in range(2):
            table.cell(r, c).text = data[r][c]
    path = tempfile.mktemp(suffix=".docx")
    document.save(path)
    return path


def test_detect_structure_docx_finds_headings_paragraphs_and_tables():
    path = _real_docx_with_heading_and_table()
    try:
        structure = detect_structure_docx(path)
        types = [e.type for e in structure]
        assert types == ["heading", "paragraph", "table"]
        assert structure[0].content == "Chapter One"
        assert "real" in structure[2].content
    finally:
        Path(path).unlink(missing_ok=True)


def test_structure_to_json_produces_a_real_recursive_shape():
    structure = [StructureElement(type="heading", level=1, content="Title", children=[
        StructureElement(type="paragraph", level=None, content="Nested"),
    ])]
    assert structure_to_json(structure) == [
        {"type": "heading", "level": 1, "content": "Title", "children": [
            {"type": "paragraph", "level": None, "content": "Nested", "children": []},
        ]},
    ]


def test_structure_to_markdown_renders_a_real_readable_document():
    """Validation criterion: la conversion en Markdown fonctionne."""
    structure = detect_structure_markdown(_MARKDOWN)
    rendered = structure_to_markdown(structure)
    assert "# Title" in rendered
    assert "## Section A" in rendered
    assert "- item one" in rendered
    assert "```" in rendered
