"""Partie 3.1.5 -- tests for api/services/image_extraction.py."""

import io

from PIL import Image

from api.services.image_extraction import extract_images_docx, extract_images_epub, extract_images_html, get_image_metadata


def _real_png_bytes(size=(10, 10), color="red") -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", size, color=color).save(buf, format="PNG")
    return buf.getvalue()


# -------------------------------------------------------------- extract_images_docx --

def test_extract_images_docx_finds_a_real_embedded_image(tmp_path):
    """Validation criterion: l'extraction d'images DOCX fonctionne."""
    import docx

    document = docx.Document()
    document.add_paragraph("before")
    document.add_picture(io.BytesIO(_real_png_bytes()))
    document.add_paragraph("after")
    path = tmp_path / "with_image.docx"
    document.save(str(path))

    images = extract_images_docx(str(path))
    assert len(images) == 1
    assert images[0].startswith(b"\x89PNG")


def test_extract_images_docx_returns_an_empty_list_for_no_images(tmp_path):
    import docx

    document = docx.Document()
    document.add_paragraph("just text, no images")
    path = tmp_path / "no_image.docx"
    document.save(str(path))

    assert extract_images_docx(str(path)) == []


# -------------------------------------------------------------- extract_images_epub --

def test_extract_images_epub_finds_a_real_embedded_image(tmp_path):
    """Validation criterion: l'extraction d'images EPUB fonctionne."""
    import ebooklib
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier("id1")
    book.set_title("Image Test Book")
    book.set_language("en")
    image_bytes = _real_png_bytes()
    img = epub.EpubItem(uid="img1", file_name="images/a.png", media_type="image/png", content=image_bytes)
    book.add_item(img)
    chapter = epub.EpubHtml(title="Chapter 1", file_name="chap1.xhtml", lang="en")
    chapter.content = "<html><body><p>Real chapter content.</p></body></html>"
    book.add_item(chapter)
    book.toc = (epub.Link("chap1.xhtml", "Chapter 1", "chap1"),)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    book.spine = ["nav", chapter]
    path = tmp_path / "with_image.epub"
    epub.write_epub(str(path), book)

    images = extract_images_epub(str(path))
    assert len(images) == 1
    assert images[0] == image_bytes


# -------------------------------------------------------------- extract_images_html --

def test_extract_images_html_finds_real_img_references(tmp_path):
    """Validation criterion: l'extraction d'images HTML fonctionne."""
    path = tmp_path / "with_images.html"
    path.write_bytes(b'<html><body><img src="https://example.com/a.png" alt="A real image"><img src="/local.jpg"></body></html>')

    images = extract_images_html(str(path))
    assert images == [
        {"src": "https://example.com/a.png", "alt": "A real image"},
        {"src": "/local.jpg", "alt": None},
    ]


def test_extract_images_html_returns_an_empty_list_for_no_images(tmp_path):
    path = tmp_path / "no_images.html"
    path.write_bytes(b"<html><body><p>No real images here.</p></body></html>")
    assert extract_images_html(str(path)) == []


# ---------------------------------------------------------------- get_image_metadata --

def test_get_image_metadata_returns_real_dimensions_and_format():
    """Validation criterion / vision critique 3: les métadonnées sont
    extraites."""
    metadata = get_image_metadata(_real_png_bytes(size=(42, 24)))
    assert metadata == {"format": "PNG", "width": 42, "height": 24}


def test_get_image_metadata_returns_empty_dict_for_real_undecodable_bytes():
    """Vision critique 3 -- que se passe-t-il si l'image est corrompue :
    jamais une exception, jamais une dimension fabriquée."""
    assert get_image_metadata(b"not a real image at all") == {}
