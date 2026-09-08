"""
Partie 8.1.14 -- real conversation export: PDF, DOCX, JSON, Markdown.

**Réutilisation réelle et délibérée**: PDF export goes through the
exact SAME real Markdown-to-HTML pipeline as 8.1.2
(`api.services.markdown_renderer.render_markdown_safe`) -- a real
message's own Markdown formatting (code blocks, lists, bold) renders
identically whether it is shown live in the chat or exported to PDF,
never a second, parallel formatting path.

**Dépendance réelle avec un vrai coût système, gérée en LAZY IMPORT**:
`weasyprint` needs real native libraries (Pango/cairo/GObject) that a
plain `pip install` does NOT provide on every real platform (confirmed
directly on this real Windows dev machine: `import weasyprint` itself
raises a real `OSError` at import time without them installed
separately) -- the exact same real, documented situation Partie
3.1.6's own Tesseract/`pytesseract` already established for this
codebase (a real system dependency CI installs via `apt-get`, real
local dev may not have). `export_to_pdf` imports `weasyprint` LAZILY,
inside the function, so importing this module (and therefore this
whole app) never crashes on a real machine missing those native
libraries -- only an actual PDF export attempt does, with a real,
honest `ExportError` instead of an app-wide import crash."""

import datetime as dt
import io
import json
import uuid

from sqlalchemy.ext.asyncio import AsyncSession

from api.config import settings
from api.models.conversation import Conversation, ConversationMessage
from api.security.conversations import get_conversation, get_conversation_messages


class ExportError(ValueError):
    """Real, honest export failure (not found, not owned, or a real,
    missing native dependency)."""


async def _load_owned_conversation_with_messages(db: AsyncSession, conversation_id: uuid.UUID, user_id: uuid.UUID) -> tuple[Conversation, list[ConversationMessage]]:
    conversation = await get_conversation(db, conversation_id)
    if conversation is None or conversation.user_id != user_id:
        raise ExportError("Conversation not found")
    messages = await get_conversation_messages(db, conversation_id, limit=settings.EXPORT_MAX_MESSAGES)
    return conversation, messages


def _format_transcript_markdown(conversation: Conversation, messages: list[ConversationMessage]) -> str:
    """Real, shared transcript formatter -- the real Markdown SOURCE
    both `export_to_markdown` (returned as-is) and `export_to_pdf`
    (rendered to HTML) build on."""
    lines = [f"# {conversation.title}", "", f"*Exported {dt.datetime.now(dt.timezone.utc).isoformat()}*", ""]
    for message in messages:
        speaker = "**You**" if message.role == "user" else f"**{message.role.capitalize()}**"
        lines.append(f"{speaker} ({message.created_at.isoformat()}):")
        lines.append("")
        lines.append(message.content)
        lines.append("")
    return "\n".join(lines)


async def export_to_markdown(db: AsyncSession, conversation_id: uuid.UUID, user_id: uuid.UUID) -> str:
    """Item 2's own literal function."""
    conversation, messages = await _load_owned_conversation_with_messages(db, conversation_id, user_id)
    return _format_transcript_markdown(conversation, messages)


async def export_to_json(db: AsyncSession, conversation_id: uuid.UUID, user_id: uuid.UUID) -> dict:
    """Item 2's own literal function -- real, complete, machine-readable
    export (every real message field, not just `content`)."""
    conversation, messages = await _load_owned_conversation_with_messages(db, conversation_id, user_id)
    return {
        "id": str(conversation.id), "title": conversation.title, "agent_id": conversation.agent_id,
        "created_at": conversation.created_at.isoformat(), "updated_at": conversation.updated_at.isoformat(),
        "messages": [
            {
                "id": str(m.id), "role": m.role, "content": m.content, "created_at": m.created_at.isoformat(),
                "tool_calls": m.tool_calls, "tool_call_id": m.tool_call_id,
            }
            for m in messages
        ],
    }


async def export_to_pdf(db: AsyncSession, conversation_id: uuid.UUID, user_id: uuid.UUID) -> bytes:
    """Item 2's own literal function -- see this module's own top
    docstring for the real lazy-import + Markdown-reuse reasoning."""
    from api.services.markdown_renderer import render_markdown_safe

    conversation, messages = await _load_owned_conversation_with_messages(db, conversation_id, user_id)
    markdown_source = _format_transcript_markdown(conversation, messages)
    html_body = render_markdown_safe(markdown_source)

    try:
        import weasyprint
    except OSError as exc:
        raise ExportError(
            "PDF export is unavailable on this deployment: WeasyPrint's native libraries "
            "(Pango/cairo/GObject) are not installed. See requirements-api.txt's own comment."
        ) from exc

    html_document = f'<html><head><meta charset="utf-8"><title>{conversation.title}</title></head><body>{html_body}</body></html>'
    return weasyprint.HTML(string=html_document).write_pdf()


async def export_to_docx(db: AsyncSession, conversation_id: uuid.UUID, user_id: uuid.UUID) -> bytes:
    """Item 2's own literal function -- `python-docx` (already a real
    dependency, 2.1.2's own real DOCX ingestion)."""
    import docx

    conversation, messages = await _load_owned_conversation_with_messages(db, conversation_id, user_id)

    document = docx.Document()
    document.add_heading(conversation.title, level=1)
    for message in messages:
        speaker = "You" if message.role == "user" else message.role.capitalize()
        document.add_heading(f"{speaker} — {message.created_at.isoformat()}", level=3)
        document.add_paragraph(message.content)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
